import csv
import io
import json
from collections import defaultdict

from django.db.models import Q
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from migracion_masiva_archivo.services.continuidad_service import analizar_continuidad
from migracion_masiva_archivo.services.error_message_service import build_phase1_error_message, build_phase2_error_message
from migracion_masiva_archivo.services.file_name_service import (
    extract_expected_consecutivo_from_filename,
    extraer_metadata_nombre_archivo,
    has_ok_marker,
    humanize_file_error,
    should_omit_system_file,
)


EXPLORATION_FIELDS = [
    'lote_id',
    'documento_id',
    'nombre_archivo',
    'ruta_archivo',
    'extension',
    'peso_bytes',
    'hash_archivo',
    'archivo_vacio',
    'archivo_sospechosamente_pequeno',
    'validacion_tamano',
    'firma_archivo',
    'extension_coincide_con_firma',
    'tipo_detectado_por_firma',
    'error_firma_archivo',
    'error_lectura_tamano',
    'error_lectura_hash',
    'archivo_posiblemente_bloqueado',
    'archivo_posiblemente_sincronizando',
    'pdf_supera_limite_saia',
    'pdf_pesado_saia',
    'advertencia_tamano_saia',
    'requiere_timeout_extendido',
    'saia_max_file_mb',
    'validacion_tamano_saia',
    'nombre_pel_ambiguo',
    'caracteres_no_compatibles_saia',
    'caracteres_problematicos_saia',
    'extension_doble_o_sospechosa',
    'duplicado_ruta_normalizada',
    'pel_simple_subparte_1_conflicto',
    'consecutivo_nombre_no_coincide_contenido',
    'consecutivos_contenido_detectados',
    'error_pdf',
    'es_soportado',
    'es_duplicado',
    'documento_duplicado_de',
    'contiene_ok',
    'consecutivo_nombre',
    'tipo_documento_nombre',
    'pel_base_nombre',
    'tomo_nombre',
    'parte_nombre',
    'subparte_nombre',
    'grupo_documental_nombre',
    'identidad_documental',
    'nombre_reconocido',
    'orden_carga',
    'orden_documental_valido',
    'motivo_orden_invalido',
    'bloqueado_por_continuidad',
    'estado_duplicado_documental',
    'duplicado_historico',
    'duplicado_historico_por',
    'documento_historico_id',
    'lote_historico_id',
    'ya_cargado_saia',
    'intento_saia_exitoso_id',
    'carga_saia_historica',
    'id_documento_saia_historico',
    'usuario_saia_historico',
    'fecha_carga_saia_historica',
    'coincidencias_historicas',
    'mensaje_bloqueo_saia_historico',
    'procesar_fase_2',
    'requiere_revision_documental',
    'observacion_duplicado',
    'codigo_error_usuario',
    'fase_error',
    'severidad_error',
    'estado_tecnico_fase_1',
    'mensaje_usuario_fase_1',
    'accion_recomendada_fase_1',
    'es_pdf_danado',
    'mensaje_usuario',
    'accion_recomendada',
    'bloquea_fase_2',
    'bloquea_saia',
    'error_tecnico',
    'fase2_codigo_error_usuario',
    'fase2_severidad_error',
    'fase2_mensaje_usuario',
    'fase2_accion_recomendada',
    'fase2_bloquea_fase_3',
    'fase2_bloquea_saia',
    'fase2_error_tecnico',
    'fase3_codigo_error_usuario',
    'fase3_severidad_error',
    'fase3_mensaje_usuario',
    'fase3_accion_recomendada',
    'fase3_bloquea_relacion',
    'fase3_bloquea_saia',
    'fase3_error_tecnico',
    'mensaje_usuario_fase_2',
    'accion_recomendada_fase_2',
    'omitido_carga_por_ok',
    'origen_archivo',
    'carpeta_origen_detectada',
    'carpeta_secundaria_usada',
    'inventario_secundaria_completo',
    'tipo_consecutivo',
    'fue_recuperado_por_continuidad',
    'consecutivo_faltante',
    'subparte_faltante',
    'tomo_faltante',
    'parte_tomo_faltante',
    'buscado_en_carpeta_secundaria',
    'encontrado_en_secundaria',
    'duplicado_entre_carpetas',
    'ruta_principal',
    'ruta_secundaria',
    'estado_continuidad',
    'estado_continuidad_base',
    'estado_continuidad_subparte',
    'estado_continuidad_tomo',
    'consecutivo_faltante_anterior',
    'observacion_continuidad',
    'procesar_saia',
    'motivo_no_procesar_saia',
    'omitido_sistema',
    'motivo_omision',
    'numero_paginas',
    'es_pdf',
    'es_digital',
    'requiere_ocr',
    'estado_proceso',
    'es_factura_principal',
    'consecutivo_normalizado',
    'nit_normalizado',
    'proveedor_detectado',
    'numeros_factura_venta_normalizados',
    'criterio_relacion',
    'confianza_relacion',
    'documento_principal',
    'documentos_relacionados',
    'nit',
    'proveedor',
    'consecutivo',
    'fecha_documento',
    'tipo_documento',
    'valor',
    'medio_pago',
    'cruce',
    'cruces_m_pago',
    'referencia_pago',
    'datos_pel',
    'pagina_detectada',
    'confianza',
    'requiere_revision',
    'observaciones',
    'error',
    'error_legible',
]


def build_exploration_csv(documents, continuity_context=None, carpeta_origen='', carpeta_secundaria=''):
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=EXPLORATION_FIELDS)
    writer.writeheader()

    for row in _documents_to_rows(documents, continuity_context, carpeta_origen, carpeta_secundaria):
        writer.writerow(row)

    return output.getvalue()


def write_exploration_csv(path, documents, continuity_context=None, carpeta_origen='', carpeta_secundaria=''):
    csv_content = build_exploration_csv(documents, continuity_context, carpeta_origen, carpeta_secundaria)
    with open(path, 'w', newline='', encoding='utf-8') as file:
        file.write(csv_content)


def build_exploration_xlsx(documents, continuity_context=None, carpeta_origen='', carpeta_secundaria=''):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Exploracion documental'

    sheet.append(EXPLORATION_FIELDS)
    header_fill = PatternFill(start_color='27348B', end_color='27348B', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row in _documents_to_rows(documents, continuity_context, carpeta_origen, carpeta_secundaria):
        sheet.append([row.get(field, '') for field in EXPLORATION_FIELDS])

    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or '')) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 60)

    _add_summary_sheet(workbook, continuity_context)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()


def write_exploration_xlsx(path, documents, continuity_context=None, carpeta_origen='', carpeta_secundaria=''):
    xlsx_content = build_exploration_xlsx(documents, continuity_context, carpeta_origen, carpeta_secundaria)
    with open(path, 'wb') as file:
        file.write(xlsx_content)


def write_exploration_report(path, documents, continuity_context=None, carpeta_origen='', carpeta_secundaria=''):
    extension = str(path).lower()
    if extension.endswith('.xlsx'):
        write_exploration_xlsx(path, documents, continuity_context, carpeta_origen, carpeta_secundaria)
        return 'xlsx'
    write_exploration_csv(path, documents, continuity_context, carpeta_origen, carpeta_secundaria)
    return 'csv'


def _document_to_row(document):
    if isinstance(document, dict):
        row = {field: document.get(field, '') for field in EXPLORATION_FIELDS}
        return _enrich_report_row(row)

    metadata = _get_document_metadata(document)
    row = {
        'lote_id': document.lote_id,
        'documento_id': document.id,
        'nombre_archivo': document.nombre_archivo,
        'ruta_archivo': document.ruta_archivo,
        'extension': document.extension,
        'peso_bytes': document.peso_bytes,
        'hash_archivo': document.hash_archivo,
        'es_soportado': document.es_soportado,
        'es_duplicado': document.es_duplicado,
        'documento_duplicado_de': document.documento_duplicado_de_id or '',
        'numero_paginas': document.numero_paginas,
        'es_pdf': document.es_pdf,
        'es_digital': document.es_digital,
        'requiere_ocr': document.requiere_ocr,
        'estado_proceso': document.estado_proceso,
        'nit': metadata.nit if metadata else '',
        'proveedor': metadata.proveedor if metadata else '',
        'consecutivo': metadata.consecutivo if metadata else '',
        'fecha_documento': metadata.fecha_documento if metadata else '',
        'tipo_documento': metadata.tipo_documento if metadata else '',
        'valor': metadata.valor if metadata else '',
        'medio_pago': metadata.medio_pago if metadata else '',
        'cruce': metadata.cruce if metadata else '',
        'cruces_m_pago': ' | '.join(str(v) for v in (metadata.cruces_m_pago or [])) if metadata else '',
        'referencia_pago': metadata.referencia_pago if metadata else '',
        'datos_pel': json.dumps(metadata.datos_pel or {}, ensure_ascii=False) if metadata else '',
        'pagina_detectada': metadata.pagina_detectada if metadata else '',
        'confianza': metadata.confianza if metadata else '',
        'requiere_revision': metadata.requiere_revision if metadata else '',
        'observaciones': metadata.observaciones if metadata else '',
        'error': document.error,
    }
    return _enrich_report_row(row)


def _documents_to_rows(documents, continuity_context=None, carpeta_origen='', carpeta_secundaria=''):
    rows = [_document_to_row(document) for document in documents]
    rows.extend(_build_continuity_alert_rows(continuity_context))
    rows = _enrich_documental_inventory(rows)
    _apply_continuity_context(rows, continuity_context, carpeta_origen, carpeta_secundaria)
    _apply_phase1_user_errors(rows)
    _apply_phase2_user_errors(rows)
    _apply_relation_context(rows)
    _apply_phase3_user_errors(rows)
    _apply_report_aliases(rows)
    return sorted(rows, key=lambda row: row.get('orden_carga') or 10**12)


def _get_document_metadata(document):
    try:
        return document.metadata
    except Exception:
        return None


def _enrich_report_row(row):
    filename = row.get('nombre_archivo', '')
    path = row.get('ruta_archivo') or filename
    error = row.get('error', '')
    omitted_system = should_omit_system_file(path)
    name_metadata = extraer_metadata_nombre_archivo(filename)

    row['contiene_ok'] = has_ok_marker(filename)
    row['consecutivo_nombre'] = extract_expected_consecutivo_from_filename(filename)
    row.update(name_metadata)
    row['omitido_sistema'] = omitted_system
    row['motivo_omision'] = 'Archivo de sistema omitido' if omitted_system else ''
    row['error_legible'] = humanize_file_error(error)
    row['omitido_carga_por_ok'] = row['contiene_ok']
    row['tipo_consecutivo'] = _get_tipo_consecutivo(row)
    row.setdefault('archivo_vacio', False)
    row.setdefault('archivo_sospechosamente_pequeno', False)
    row.setdefault('validacion_tamano', '')
    row.setdefault('firma_archivo', '')
    row.setdefault('extension_coincide_con_firma', '')
    row.setdefault('tipo_detectado_por_firma', '')
    row.setdefault('error_firma_archivo', '')
    row.setdefault('error_lectura_tamano', '')
    row.setdefault('error_lectura_hash', '')
    row.setdefault('archivo_posiblemente_bloqueado', False)
    row.setdefault('archivo_posiblemente_sincronizando', False)
    row.setdefault('pdf_supera_limite_saia', False)
    row.setdefault('pdf_pesado_saia', row.get('pdf_supera_limite_saia', False))
    row.setdefault('advertencia_tamano_saia', '')
    row.setdefault('requiere_timeout_extendido', row.get('pdf_pesado_saia', False))
    row.setdefault('saia_max_file_mb', '')
    row.setdefault('validacion_tamano_saia', '')
    row.setdefault('nombre_pel_ambiguo', False)
    row.setdefault('caracteres_no_compatibles_saia', False)
    row.setdefault('caracteres_problematicos_saia', '')
    row.setdefault('extension_doble_o_sospechosa', False)
    row.setdefault('duplicado_ruta_normalizada', False)
    row.setdefault('pel_simple_subparte_1_conflicto', False)
    row.setdefault('consecutivo_nombre_no_coincide_contenido', False)
    row.setdefault('consecutivos_contenido_detectados', '')
    row.setdefault('error_pdf', '')
    row.setdefault('orden_documental_valido', True)
    row.setdefault('motivo_orden_invalido', '')
    row.setdefault('bloqueado_por_continuidad', False)
    row.setdefault('duplicado_historico', False)
    row.setdefault('duplicado_historico_por', '')
    row.setdefault('documento_historico_id', '')
    row.setdefault('lote_historico_id', '')
    row.setdefault('ya_cargado_saia', False)
    row.setdefault('intento_saia_exitoso_id', '')
    row.setdefault('carga_saia_historica', False)
    row.setdefault('id_documento_saia_historico', '')
    row.setdefault('usuario_saia_historico', '')
    row.setdefault('fecha_carga_saia_historica', '')
    row.setdefault('coincidencias_historicas', '')
    row.setdefault('mensaje_bloqueo_saia_historico', '')
    row.setdefault('codigo_error_usuario', '')
    row.setdefault('fase_error', '')
    row.setdefault('severidad_error', '')
    row.setdefault('estado_tecnico_fase_1', '')
    row.setdefault('mensaje_usuario_fase_1', '')
    row.setdefault('accion_recomendada_fase_1', '')
    row.setdefault('es_pdf_danado', False)
    row.setdefault('mensaje_usuario', '')
    row.setdefault('accion_recomendada', '')
    row.setdefault('bloquea_fase_2', False)
    row.setdefault('bloquea_saia', False)
    row.setdefault('error_tecnico', '')
    row.setdefault('fase2_codigo_error_usuario', '')
    row.setdefault('fase2_severidad_error', '')
    row.setdefault('fase2_mensaje_usuario', '')
    row.setdefault('fase2_accion_recomendada', '')
    row.setdefault('fase2_bloquea_fase_3', False)
    row.setdefault('fase2_bloquea_saia', False)
    row.setdefault('fase2_error_tecnico', '')
    row.setdefault('fase3_codigo_error_usuario', '')
    row.setdefault('fase3_severidad_error', '')
    row.setdefault('fase3_mensaje_usuario', '')
    row.setdefault('fase3_accion_recomendada', '')
    row.setdefault('fase3_bloquea_relacion', False)
    row.setdefault('fase3_bloquea_saia', False)
    row.setdefault('fase3_error_tecnico', '')
    row.setdefault('mensaje_usuario_fase_2', '')
    row.setdefault('accion_recomendada_fase_2', '')
    row.setdefault('origen_archivo', '')
    row.setdefault('carpeta_origen_detectada', '')
    row.setdefault('carpeta_secundaria_usada', False)
    row.setdefault('inventario_secundaria_completo', False)
    row.setdefault('fue_recuperado_por_continuidad', False)
    row.setdefault('consecutivo_faltante', '')
    row.setdefault('subparte_faltante', '')
    row.setdefault('tomo_faltante', '')
    row.setdefault('parte_tomo_faltante', '')
    row.setdefault('buscado_en_carpeta_secundaria', False)
    row.setdefault('encontrado_en_secundaria', False)
    row.setdefault('duplicado_entre_carpetas', False)
    row.setdefault('ruta_principal', '')
    row.setdefault('ruta_secundaria', '')
    row.setdefault('estado_continuidad', '')
    row.setdefault('estado_continuidad_base', '')
    row.setdefault('estado_continuidad_subparte', '')
    row.setdefault('estado_continuidad_tomo', '')
    row.setdefault('procesar_saia', False)
    row.setdefault('motivo_no_procesar_saia', '')
    row.setdefault('es_factura_principal', '')
    row.setdefault('consecutivo_normalizado', '')
    row.setdefault('nit_normalizado', '')
    row.setdefault('proveedor_detectado', '')
    row.setdefault('numeros_factura_venta_normalizados', '')
    row.setdefault('criterio_relacion', '')
    row.setdefault('confianza_relacion', '')
    row.setdefault('documento_principal', '')
    row.setdefault('documentos_relacionados', '')
    return row


def _enrich_documental_inventory(rows):
    identity_groups = defaultdict(list)
    hash_groups = defaultdict(list)
    base_groups = defaultdict(list)

    for row in rows:
        identity = row.get('identidad_documental') or ''
        file_hash = row.get('hash_archivo') or ''
        base = row.get('pel_base_nombre') or ''

        if identity:
            identity_groups[identity].append(row)
        if file_hash:
            hash_groups[file_hash].append(row)
        if base:
            base_groups[base].append(row)

    for row in rows:
        _classify_documental_row(row, identity_groups, hash_groups, base_groups)

    _assign_load_order(rows)
    _assign_continuity_observations(rows)
    for row in rows:
        _assign_saia_processing(row)
    return rows


def _classify_documental_row(row, identity_groups, hash_groups, base_groups):
    identity = row.get('identidad_documental') or ''
    file_hash = row.get('hash_archivo') or ''
    base = row.get('pel_base_nombre') or ''
    has_error = bool(row.get('error'))
    is_supported = bool(row.get('es_soportado'))

    row['estado_duplicado_documental'] = 'UNICO'
    row['procesar_fase_2'] = True
    row['requiere_revision_documental'] = False
    row['observacion_duplicado'] = ''

    if not row.get('nombre_reconocido'):
        row['estado_duplicado_documental'] = 'NOMBRE_NO_RECONOCIDO'
        row['procesar_fase_2'] = False
        row['requiere_revision_documental'] = True
        row['observacion_duplicado'] = 'No se reconocio el PEL desde el nombre del archivo'
        return

    if identity and len(identity_groups[identity]) > 1:
        row['estado_duplicado_documental'] = 'DUPLICADO_REAL'
        row['procesar_fase_2'] = False
        row['requiere_revision_documental'] = True
        row['observacion_duplicado'] = 'Existe otro archivo con la misma identidad documental'
        return

    if file_hash and len(hash_groups[file_hash]) > 1:
        row['estado_duplicado_documental'] = 'POSIBLE_DUPLICADO_HASH'
        row['observacion_duplicado'] = (
            'Comparte hash con otro archivo, pero tiene identidad documental diferente'
        )
    elif base and len(base_groups[base]) > 1:
        row['estado_duplicado_documental'] = 'MULTIPARTE_RELACIONADO'
        row['observacion_duplicado'] = 'Comparte PEL base con otras subpartes o tomos'

    if not is_supported:
        row['procesar_fase_2'] = False
        row['requiere_revision_documental'] = True
        row['observacion_duplicado'] = _append_observation(
            row['observacion_duplicado'],
            'Formato no soportado para Fase 2',
        )
    elif (
        has_error
        and row.get('estado_duplicado_documental') != 'POSIBLE_DUPLICADO_HASH'
        and _row_error_blocks_phase2(row)
    ):
        row['procesar_fase_2'] = False
        row['requiere_revision_documental'] = True
        row['observacion_duplicado'] = _append_observation(
            row['observacion_duplicado'],
            'El archivo tiene error tecnico y requiere revision',
        )

    if row.get('nombre_pel_ambiguo') or row.get('extension_doble_o_sospechosa'):
        row['procesar_fase_2'] = False
        row['requiere_revision_documental'] = True
    if row.get('pel_simple_subparte_1_conflicto') or row.get('consecutivo_nombre_no_coincide_contenido'):
        row['procesar_fase_2'] = False
        row['requiere_revision_documental'] = True
    if 'duplicado por ruta normalizada' in str(row.get('error', '')).lower():
        row['duplicado_ruta_normalizada'] = True


def _row_error_blocks_phase2(row):
    if row.get('nombre_pel_ambiguo') or row.get('extension_doble_o_sospechosa'):
        return True
    if row.get('consecutivo_nombre_no_coincide_contenido'):
        return True
    if row.get('archivo_vacio') or row.get('archivo_sospechosamente_pequeno'):
        return True
    if row.get('error_lectura_tamano') or row.get('error_lectura_hash'):
        return True
    if row.get('error_firma_archivo') or row.get('error_pdf'):
        return True
    if row.get('es_soportado') and row.get('extension_coincide_con_firma') is False:
        return True
    text = str(row.get('error') or '').lower()
    return any(
        marker in text
        for marker in (
            'formato no soportado',
            'extension no coincide',
            'archivo vacio',
            'sospechosamente pequeno',
            'pdf posiblemente',
            'pdf no se pudo abrir',
        )
    )


def _assign_load_order(rows):
    ordered_rows = sorted(enumerate(rows), key=lambda item: _load_order_key(item[1], item[0]))
    for order, (_, row) in enumerate(ordered_rows, start=1):
        row['orden_carga'] = order


def _load_order_key(row, original_index):
    base = _safe_int(row.get('pel_base_nombre'))
    tomo = _safe_int(row.get('tomo_nombre'))
    parte = _safe_int(row.get('parte_nombre'))
    subparte = _safe_int(row.get('subparte_nombre')) or 1

    if base is None:
        return (1, 10**12, 10**12, 10**12, 10**12, original_index)

    if tomo is not None:
        return (0, base, tomo, parte or 1, 0, original_index)

    return (0, base, 0, subparte, 0, original_index)


def _assign_continuity_observations(rows):
    bases = sorted({_safe_int(row.get('pel_base_nombre')) for row in rows if _safe_int(row.get('pel_base_nombre'))})
    missing_before_base = {}
    previous = None

    for base in bases:
        if previous is not None and base > previous + 1:
            gap_size = base - previous - 1
            if gap_size <= 1000:
                missing_before_base[base] = _format_missing_pel_range(previous + 1, base - 1)
        previous = base

    for row in rows:
        if row.get('estado_continuidad') in {
            'FALTANTE_NO_ENCONTRADO',
            'CARPETA_SECUNDARIA_NO_CONFIGURADA',
            'CARPETA_SECUNDARIA_NO_ACCESIBLE',
        }:
            continue
        base = _safe_int(row.get('pel_base_nombre'))
        missing = missing_before_base.get(base, '')
        row['consecutivo_faltante_anterior'] = missing
        row['observacion_continuidad'] = (
            f'Faltan consecutivos antes de PEL {base}: {missing}' if missing else ''
        )


def _build_continuity_alert_rows(context):
    rows = []
    if not context:
        return rows

    for alert in context.get('alertas', []):
        rows.append(
            _build_alert_row(
                alert.get('nombre_archivo', ''),
                alert.get('estado_continuidad', ''),
                alert.get('observacion_continuidad', ''),
                consecutivo_faltante=alert.get('consecutivo_faltante', ''),
                buscado_en_carpeta_secundaria=alert.get('estado_continuidad')
                != 'CARPETA_SECUNDARIA_NO_CONFIGURADA',
            )
        )

    for alert in context.get('subpartes', {}).get('faltantes', []):
        rows.append(
            _build_alert_row(
                alert.get('consecutivo_faltante', ''),
                alert.get('estado_continuidad', ''),
                f"Falta subparte documental: {alert.get('consecutivo_faltante')}",
                subparte_faltante=alert.get('consecutivo_faltante', ''),
                estado_continuidad_subparte=alert.get('estado_continuidad', ''),
            )
        )

    for alert in context.get('tomos', {}).get('faltantes', []):
        estado_tomo = alert.get('estado_continuidad', '')
        tomo_faltante = ''
        parte_tomo_faltante = ''
        if estado_tomo in {'FALTA_TOMO_INICIAL', 'FALTANTE_TOMO'}:
            tomo_faltante = alert.get('consecutivo_faltante', '')
        if estado_tomo in {'FALTA_PARTE_INICIAL', 'FALTANTE_PARTE_TOMO', 'FALTANTE_PARTE_TOMO_NO_ENCONTRADO'}:
            parte_tomo_faltante = alert.get('consecutivo_faltante', '')
        rows.append(
            _build_alert_row(
                alert.get('consecutivo_faltante', ''),
                estado_tomo,
                f"Falta tomo o parte documental: {alert.get('consecutivo_faltante')}",
                tomo_faltante=tomo_faltante,
                parte_tomo_faltante=parte_tomo_faltante,
                estado_continuidad_tomo=estado_tomo,
            )
        )
    return rows


def _apply_continuity_context(rows, context, carpeta_origen='', carpeta_secundaria=''):
    primary_root = _normalize_path(carpeta_origen)
    secondary_root = _normalize_path(carpeta_secundaria)
    resumen = (context or {}).get('resumen', {})
    recovered_by_path = {
        _normalize_path(item.get('ruta_archivo')): item for item in (context or {}).get('recuperados', [])
    }
    duplicate_by_primary_path = {
        _normalize_path(item.get('ruta_principal')): item
        for item in (context or {}).get('duplicados_entre_carpetas', [])
    }
    changes_by_base = {
        _safe_int(item.get('hasta')): item for item in (context or {}).get('cambios_rango', [])
    }
    validation_by_document_id = (context or {}).get('validaciones_archivo', {})
    historic_duplicates_by_document_id = {
        item.get('documento_id'): item for item in (context or {}).get('duplicados_historicos', [])
    }
    simple_subpart_conflict_ids = {
        doc_id
        for conflict in (context or {}).get('conflictos_simple_subparte_1', [])
        for doc_id in conflict.get('documentos', [])
    }
    blocked_bases = _continuity_blocked_bases(context)

    for row in rows:
        row_path = _normalize_path(row.get('ruta_archivo'))
        base = _safe_int(row.get('pel_base_nombre'))
        document_id = row.get('documento_id')
        if document_id in validation_by_document_id:
            row.update(validation_by_document_id[document_id])
        if document_id in historic_duplicates_by_document_id:
            duplicate = historic_duplicates_by_document_id[document_id]
            row['duplicado_historico'] = True
            row['duplicado_historico_por'] = duplicate.get('duplicado_historico_por', '')
            row['documento_historico_id'] = duplicate.get('documento_historico_id', '')
            row['lote_historico_id'] = duplicate.get('lote_historico_id', '')
            row['ya_cargado_saia'] = duplicate.get('ya_cargado_saia', False)
            row['intento_saia_exitoso_id'] = duplicate.get('intento_saia_exitoso_id', '')
            row['carga_saia_historica'] = duplicate.get('carga_saia_historica', False)
            row['id_documento_saia_historico'] = duplicate.get('id_documento_saia_historico', '')
            row['usuario_saia_historico'] = duplicate.get('usuario_saia_historico', '')
            row['fecha_carga_saia_historica'] = duplicate.get('fecha_carga_saia_historica', '')
            row['coincidencias_historicas'] = duplicate.get('coincidencias_historicas', '')
            row['mensaje_bloqueo_saia_historico'] = duplicate.get('mensaje_bloqueo_saia_historico', '')
            row['estado_proceso'] = 'REQUIERE_REVISION'
            row['requiere_revision_documental'] = True
            row['procesar_saia'] = False
            row['motivo_no_procesar_saia'] = (
                row['mensaje_bloqueo_saia_historico']
                or 'Duplicado historico pendiente de decision documental'
            )
        if document_id in simple_subpart_conflict_ids:
            row['pel_simple_subparte_1_conflicto'] = True
            row['estado_continuidad'] = 'PEL_SIMPLE_Y_SUBPARTE_1_CONFLICTO'
            row['observacion_continuidad'] = 'Conflicto entre PEL simple y subparte 1'
            row['procesar_fase_2'] = False
            row['requiere_revision_documental'] = True
            row['procesar_saia'] = False
            row['motivo_no_procesar_saia'] = 'Conflicto entre PEL simple y subparte 1'
        row['carpeta_secundaria_usada'] = resumen.get('carpeta_secundaria_usada', False)
        row['inventario_secundaria_completo'] = resumen.get('inventario_secundaria_completo', False)

        if row.get('estado_continuidad') in {
            'FALTANTE_NO_ENCONTRADO',
            'CARPETA_SECUNDARIA_NO_CONFIGURADA',
            'CARPETA_SECUNDARIA_NO_ACCESIBLE',
            'FALTANTE_SUBPARTE',
            'FALTA_PRINCIPAL_SUBPARTE',
            'FALTANTE_TOMO',
            'FALTANTE_PARTE_TOMO',
            'FALTA_TOMO_INICIAL',
            'FALTA_PARTE_INICIAL',
            'FALTANTE_SUBPARTE_NO_ENCONTRADO',
            'FALTANTE_PARTE_TOMO_NO_ENCONTRADO',
        }:
            row['origen_archivo'] = 'SIN_ARCHIVO'
            row['carpeta_origen_detectada'] = ''
            row['fue_recuperado_por_continuidad'] = False
            row['orden_documental_valido'] = False
            row['motivo_orden_invalido'] = row.get('observacion_continuidad') or row.get('error') or 'Alerta de continuidad'
            row['bloqueado_por_continuidad'] = True
            row['procesar_fase_2'] = False
            row['requiere_revision_documental'] = True
            continue

        if row_path in recovered_by_path:
            recovered = recovered_by_path[row_path]
            status = recovered.get('estado_continuidad') or 'INVENTARIO_SECUNDARIA_COMPLETO'
            was_missing_recovered = status == 'FALTANTE_ENCONTRADO_EN_SECUNDARIA'
            if status in {
                'FALTANTE_SUBPARTE_ENCONTRADO_EN_SECUNDARIA',
                'FALTANTE_PARTE_TOMO_ENCONTRADO_EN_SECUNDARIA',
            }:
                was_missing_recovered = True
            row['origen_archivo'] = 'SECUNDARIA_245'
            row['carpeta_origen_detectada'] = carpeta_secundaria
            row['fue_recuperado_por_continuidad'] = was_missing_recovered
            row['consecutivo_faltante'] = f"PEL {recovered.get('pel_base')}" if was_missing_recovered else ''
            if recovered.get('consecutivo_faltante'):
                row['consecutivo_faltante'] = recovered.get('consecutivo_faltante')
            row['buscado_en_carpeta_secundaria'] = was_missing_recovered
            row['encontrado_en_secundaria'] = was_missing_recovered
            row['ruta_secundaria'] = recovered.get('ruta_archivo', '')
            row['estado_continuidad'] = status
            row['estado_continuidad_base'] = status
            row['observacion_continuidad'] = recovered.get('observacion_continuidad', '')
            _assign_saia_processing(row)
            continue

        if row_path in duplicate_by_primary_path:
            duplicate = duplicate_by_primary_path[row_path]
            row['duplicado_entre_carpetas'] = True
            row['ruta_principal'] = duplicate.get('ruta_principal', '')
            row['ruta_secundaria'] = duplicate.get('ruta_secundaria', '')
            row['estado_continuidad'] = 'DUPLICADO_ENTRE_CARPETAS'
            row['observacion_continuidad'] = 'Documento tambien existe en carpeta secundaria; se prefiere principal'

        if secondary_root and row_path.startswith(secondary_root):
            row['origen_archivo'] = 'SECUNDARIA_245'
            row['carpeta_origen_detectada'] = carpeta_secundaria
        elif primary_root and row_path.startswith(primary_root):
            row['origen_archivo'] = 'PRINCIPAL_246'
            row['carpeta_origen_detectada'] = carpeta_origen

        if base in changes_by_base:
            change = changes_by_base[base]
            row['estado_continuidad'] = 'CAMBIO_DE_RANGO'
            row['estado_continuidad_base'] = 'CAMBIO_DE_RANGO'
            row['observacion_continuidad'] = (
                f"Cambio de rango entre PEL {change.get('desde')} y PEL {change.get('hasta')}; "
                f"se omitieron {change.get('faltantes_omitidos')} consecutivos"
            )
        elif not row.get('estado_continuidad'):
            row['estado_continuidad'] = 'SECUENCIA_OK'
            row['estado_continuidad_base'] = 'SECUENCIA_OK'
        if not row.get('estado_continuidad_subparte'):
            row['estado_continuidad_subparte'] = 'SECUENCIA_OK'
        if not row.get('estado_continuidad_tomo'):
            row['estado_continuidad_tomo'] = (
                'SECUENCIA_TOMO_OK' if row.get('tomo_nombre') else ''
            )

        if base in blocked_bases:
            row['orden_documental_valido'] = False
            row['bloqueado_por_continuidad'] = True
            row['motivo_orden_invalido'] = blocked_bases[base]
            row['procesar_fase_2'] = False
            row['requiere_revision_documental'] = True

        _assign_saia_processing(row)

    if not context:
        _apply_local_continuity_without_context(rows)


def _apply_local_continuity_without_context(rows):
    analysis = analizar_continuidad(rows)
    missing_by_next = defaultdict(list)
    for item in analysis['faltantes']:
        missing_by_next[item['siguiente']].append(f"PEL {item['pel_base']}")

    for row in rows:
        base = _safe_int(row.get('pel_base_nombre'))
        if base in missing_by_next:
            missing = ', '.join(missing_by_next[base])
            row['consecutivo_faltante_anterior'] = missing
            row['observacion_continuidad'] = f'Faltan consecutivos antes de PEL {base}: {missing}'


def _build_alert_row(
    nombre_archivo,
    estado_continuidad,
    observacion,
    consecutivo_faltante='',
    subparte_faltante='',
    tomo_faltante='',
    parte_tomo_faltante='',
    estado_continuidad_subparte='',
    estado_continuidad_tomo='',
    buscado_en_carpeta_secundaria=False,
):
    row = {field: '' for field in EXPLORATION_FIELDS}
    row.update(
        {
            'nombre_archivo': nombre_archivo,
            'ruta_archivo': '',
            'extension': '',
            'es_soportado': False,
            'estado_proceso': 'REQUIERE_REVISION',
            'error': observacion,
        }
    )
    row = _enrich_report_row(row)
    row.update(
        {
            'consecutivo_faltante': consecutivo_faltante,
            'subparte_faltante': subparte_faltante,
            'tomo_faltante': tomo_faltante,
            'parte_tomo_faltante': parte_tomo_faltante,
            'buscado_en_carpeta_secundaria': buscado_en_carpeta_secundaria,
            'encontrado_en_secundaria': False,
            'estado_continuidad': estado_continuidad,
            'estado_continuidad_base': (
                estado_continuidad if estado_continuidad in {'FALTANTE_PEL_BASE', 'CAMBIO_DE_RANGO'} else ''
            ),
            'estado_continuidad_subparte': estado_continuidad_subparte,
            'estado_continuidad_tomo': estado_continuidad_tomo,
            'observacion_continuidad': observacion,
            'procesar_fase_2': False,
            'procesar_saia': False,
            'motivo_no_procesar_saia': 'Alerta de continuidad sin archivo fisico',
            'requiere_revision_documental': True,
        }
    )
    return row


def _get_tipo_consecutivo(row):
    if not row.get('nombre_reconocido'):
        return 'NO_RECONOCIDO'
    if row.get('tomo_nombre'):
        return 'PEL_TOMO_PARTE'
    if row.get('subparte_nombre') and int(row.get('subparte_nombre') or 0) > 1:
        return 'PEL_SUBPARTE'
    return 'PEL_SIMPLE'


def _assign_saia_processing(row):
    if row.get('contiene_ok'):
        row['procesar_saia'] = False
        row['motivo_no_procesar_saia'] = 'Archivo ya contiene OK'
        return
    if row.get('estado_proceso') == 'REQUIERE_REVISION':
        row['procesar_saia'] = False
        row['motivo_no_procesar_saia'] = (
            row.get('mensaje_bloqueo_saia_historico')
            or row.get('mensaje_usuario')
            or row.get('error')
            or 'Documento requiere revision documental'
        )
        return
    if row.get('estado_continuidad') in {
        'FALTANTE_NO_ENCONTRADO',
        'CARPETA_SECUNDARIA_NO_CONFIGURADA',
        'CARPETA_SECUNDARIA_NO_ACCESIBLE',
        'FALTANTE_SUBPARTE',
        'FALTA_PRINCIPAL_SUBPARTE',
        'FALTANTE_TOMO',
        'FALTANTE_PARTE_TOMO',
        'FALTA_TOMO_INICIAL',
        'FALTA_PARTE_INICIAL',
        'FALTANTE_SUBPARTE_NO_ENCONTRADO',
        'FALTANTE_PARTE_TOMO_NO_ENCONTRADO',
        'PEL_SIMPLE_Y_SUBPARTE_1_CONFLICTO',
    }:
        row['procesar_saia'] = False
        row['motivo_no_procesar_saia'] = 'Requiere revision de continuidad'
        return
    if row.get('bloqueado_por_continuidad'):
        row['procesar_saia'] = False
        row['motivo_no_procesar_saia'] = 'Requiere revision de continuidad'
        return
    if row.get('ya_cargado_saia'):
        row['procesar_saia'] = False
        row['motivo_no_procesar_saia'] = (
            row.get('mensaje_bloqueo_saia_historico')
            or 'Documento ya cargado exitosamente a SAIA'
        )
        return
    if row.get('caracteres_no_compatibles_saia'):
        row['procesar_saia'] = False
        row['motivo_no_procesar_saia'] = 'Nombre con caracteres problematicos para SAIA'
        return
    if row.get('duplicado_ruta_normalizada'):
        row['procesar_saia'] = False
        row['motivo_no_procesar_saia'] = 'Duplicado por ruta normalizada'
        return
    if row.get('pel_simple_subparte_1_conflicto'):
        row['procesar_saia'] = False
        row['motivo_no_procesar_saia'] = 'Conflicto entre PEL simple y subparte 1'
        return
    if not row.get('procesar_fase_2'):
        row['procesar_saia'] = False
        row['motivo_no_procesar_saia'] = 'No esta habilitado para Fase 2'
        return
    row['procesar_saia'] = True
    row['motivo_no_procesar_saia'] = ''


def _apply_phase1_user_errors(rows):
    for row in rows:
        error_message = build_phase1_error_message(_row_technical_error(row), row)
        row.update(error_message)

        if error_message['bloquea_fase_2']:
            row['procesar_fase_2'] = False
            row['requiere_revision_documental'] = True
        if error_message['bloquea_saia']:
            row['procesar_saia'] = False
            if not row.get('motivo_no_procesar_saia'):
                row['motivo_no_procesar_saia'] = error_message['mensaje_usuario']
        if error_message.get('codigo_error_usuario') == 'DUPLICADO_HISTORICO':
            row['procesar_saia'] = False
            row['requiere_revision_documental'] = True
            if not row.get('motivo_no_procesar_saia'):
                row['motivo_no_procesar_saia'] = (
                    row.get('mensaje_bloqueo_saia_historico')
                    or row.get('error')
                    or 'Duplicado historico pendiente de decision documental'
                )


def _apply_phase2_user_errors(rows):
    for row in rows:
        error_message = build_phase2_error_message(_row_phase2_technical_error(row), row)
        if not error_message.get('codigo_error_usuario'):
            continue

        row['fase2_codigo_error_usuario'] = error_message['codigo_error_usuario']
        row['fase2_severidad_error'] = error_message['severidad_error']
        row['fase2_mensaje_usuario'] = error_message['mensaje_usuario']
        row['fase2_accion_recomendada'] = error_message['accion_recomendada']
        row['fase2_bloquea_fase_3'] = error_message['bloquea_fase_3']
        row['fase2_bloquea_saia'] = error_message['bloquea_saia']
        row['fase2_error_tecnico'] = error_message['error_tecnico']

        if error_message['bloquea_fase_3']:
            row['requiere_revision_documental'] = True
        if error_message['bloquea_saia']:
            row['procesar_saia'] = False
            if not row.get('motivo_no_procesar_saia'):
                row['motivo_no_procesar_saia'] = error_message['mensaje_usuario']


def _row_technical_error(row):
    parts = [
        row.get('error'),
        row.get('observacion_continuidad'),
        row.get('observacion_duplicado'),
        row.get('motivo_orden_invalido'),
        row.get('motivo_no_procesar_saia'),
    ]
    return '; '.join(str(part) for part in parts if part)


def _row_phase2_technical_error(row):
    if not _row_has_phase2_context(row):
        return ''
    parts = [
        row.get('observaciones'),
        row.get('error') if row.get('estado_proceso') == 'REQUIERE_REVISION' else '',
    ]
    return '; '.join(str(part) for part in parts if part)


def _row_has_phase2_context(row):
    metadata_fields = [
        'nit',
        'proveedor',
        'consecutivo',
        'tipo_documento',
        'cruce',
        'cruces_m_pago',
        'datos_pel',
        'confianza',
        'requiere_revision',
        'observaciones',
    ]
    return any(row.get(field) not in ('', None, False, []) for field in metadata_fields)


def _apply_relation_context(rows):
    row_by_id = {}
    ids = []
    for row in rows:
        document_id = _safe_int(row.get('documento_id'))
        if document_id is None:
            continue
        row_by_id[document_id] = row
        ids.append(document_id)

    if not ids:
        return

    from migracion_masiva_archivo.models import RelacionDocumento

    relations = RelacionDocumento.objects.filter(
        Q(documento_principal_id__in=ids) | Q(documento_relacionado_id__in=ids)
    ).select_related('documento_principal', 'documento_relacionado')

    related_names = defaultdict(list)
    relation_criteria = defaultdict(list)
    relation_confidence = defaultdict(list)
    principal_names = defaultdict(list)

    for relation in relations:
        principal_id = relation.documento_principal_id
        related_id = relation.documento_relacionado_id
        related_label = _document_label(relation.documento_relacionado)
        principal_label = _document_label(relation.documento_principal)
        criterio = relation.criterio_relacion or ''
        confianza = str(relation.confianza)

        related_names[principal_id].append(related_label)
        relation_criteria[principal_id].append(criterio)
        relation_confidence[principal_id].append(confianza)

        principal_names[related_id].append(principal_label)
        relation_criteria[related_id].append(criterio)
        relation_confidence[related_id].append(confianza)

    for document_id, row in row_by_id.items():
        row['documentos_relacionados'] = ' | '.join(_unique_keep_order(related_names[document_id]))
        row['documento_principal'] = ' | '.join(_unique_keep_order(principal_names[document_id]))
        row['criterio_relacion'] = ' | '.join(_unique_keep_order(relation_criteria[document_id]))
        row['confianza_relacion'] = ' | '.join(_unique_keep_order(relation_confidence[document_id]))


def _apply_phase3_user_errors(rows):
    row_by_id = {
        _safe_int(row.get('documento_id')): row
        for row in rows
        if _safe_int(row.get('documento_id')) is not None
    }
    if not row_by_id:
        return

    lote_ids = {
        _safe_int(row.get('lote_id'))
        for row in rows
        if _safe_int(row.get('lote_id')) is not None
    }
    if not lote_ids:
        return

    from migracion_masiva_archivo.models import LogProcesoDocumental

    logs = LogProcesoDocumental.objects.filter(
        lote_id__in=lote_ids,
        evento='relacion_documental',
    ).order_by('-creado')

    applied = set()
    for log in logs:
        detail = log.detalle or {}
        for collection_name in [
            'grupos_ambiguos',
            'relaciones_descartadas',
            'documentos_omitidos',
            'advertencias_orden',
        ]:
            for item in detail.get(collection_name, []) or []:
                for document_id in _phase3_item_document_ids(item):
                    if document_id in applied:
                        continue
                    row = row_by_id.get(document_id)
                    if not row:
                        continue
                    row['fase3_codigo_error_usuario'] = item.get('codigo_error_usuario', '')
                    row['fase3_severidad_error'] = item.get('severidad_error', '')
                    row['fase3_mensaje_usuario'] = item.get('mensaje_usuario', '')
                    row['fase3_accion_recomendada'] = item.get('accion_recomendada', '')
                    row['fase3_bloquea_relacion'] = item.get('bloquea_relacion', False)
                    row['fase3_bloquea_saia'] = item.get('bloquea_saia', False)
                    row['fase3_error_tecnico'] = item.get('error_tecnico', '')
                    applied.add(document_id)


def _phase3_item_document_ids(item):
    ids = []
    for key in ['documento', 'documento_principal', 'documento_relacionado']:
        value = _safe_int(item.get(key))
        if value is not None:
            ids.append(value)
    for value in item.get('documentos', []) or []:
        doc_id = _safe_int(value)
        if doc_id is not None:
            ids.append(doc_id)
    return _unique_keep_order(ids)


def _apply_report_aliases(rows):
    for row in rows:
        datos_pel = _datos_pel_dict(row.get('datos_pel'))
        row['estado_tecnico_fase_1'] = _phase1_technical_status(row)
        row['mensaje_usuario_fase_1'] = row.get('mensaje_usuario') or _phase1_success_message(row)
        row['accion_recomendada_fase_1'] = row.get('accion_recomendada') or _phase1_success_action(row)
        row['es_pdf_danado'] = _is_damaged_pdf_row(row)

        row['es_factura_principal'] = datos_pel.get('es_factura_principal', '')
        row['consecutivo_normalizado'] = row.get('consecutivo') or datos_pel.get('consecutivo', '')
        row['nit_normalizado'] = row.get('nit', '')
        row['proveedor_detectado'] = row.get('proveedor', '')
        row['numeros_factura_venta_normalizados'] = ' | '.join(
            str(value) for value in (datos_pel.get('numeros_factura_venta_normalizados') or [])
        )
        row['mensaje_usuario_fase_2'] = row.get('fase2_mensaje_usuario', '')
        row['accion_recomendada_fase_2'] = row.get('fase2_accion_recomendada', '')


def _phase1_technical_status(row):
    code = row.get('codigo_error_usuario') or ''
    if row.get('ya_cargado_saia') or code in {'YA_CARGADO_SAIA', 'YA_CARGADO_SAIA_HISTORICO'}:
        return 'YA_CARGADO_SAIA'
    if row.get('contiene_ok') or code == 'OK_DETECTADO':
        return 'YA_TIENE_OK'
    if row.get('omitido_sistema'):
        return 'OMITIDO_SISTEMA'
    if row.get('estado_continuidad') in {
        'FALTANTE_NO_ENCONTRADO',
        'FALTANTE_PEL_BASE',
        'FALTANTE_SUBPARTE_NO_ENCONTRADO',
        'FALTANTE_PARTE_TOMO_NO_ENCONTRADO',
        'CARPETA_SECUNDARIA_NO_CONFIGURADA',
        'CARPETA_SECUNDARIA_NO_ACCESIBLE',
    }:
        return 'FALTANTE_CONTINUIDAD'
    if code in {'PDF_DANADO', 'PDF_CORRUPTO', 'ARCHIVO_VACIO', 'EXTENSION_NO_COINCIDE', 'EXTENSION_DOBLE_O_SOSPECHOSA'}:
        return 'FALLIDO_TECNICO'
    if row.get('duplicado_historico') or code == 'DUPLICADO_HISTORICO':
        return 'DUPLICADO_HISTORICO'
    if code or row.get('requiere_revision_documental') or row.get('error'):
        return 'REQUIERE_REVISION'
    return 'INVENTARIADO_OK'


def _phase1_success_message(row):
    if row.get('estado_tecnico_fase_1') == 'INVENTARIADO_OK':
        return 'El archivo fue inventariado correctamente.'
    return ''


def _phase1_success_action(row):
    if row.get('estado_tecnico_fase_1') == 'INVENTARIADO_OK':
        return 'Puede continuar con Fase 2 si el documento no tiene OK y no esta bloqueado por continuidad.'
    return ''


def _is_damaged_pdf_row(row):
    text = ' '.join(
        str(row.get(field) or '').lower()
        for field in ['codigo_error_usuario', 'error', 'error_legible', 'mensaje_usuario']
    )
    return any(pattern in text for pattern in ['pdf_danado', 'pdf_corrupto', 'eof marker', 'stream has ended', 'broken document'])


def _datos_pel_dict(value):
    if isinstance(value, dict):
        return value
    if not value:
        return {}
    try:
        loaded = json.loads(value)
    except (TypeError, ValueError):
        return {}
    return loaded if isinstance(loaded, dict) else {}


def _document_label(document):
    return f'{document.id} - {document.nombre_archivo}'


def _unique_keep_order(values):
    result = []
    seen = set()
    for value in values:
        if not value or value in seen:
            continue
        result.append(value)
        seen.add(value)
    return result


def _add_summary_sheet(workbook, continuity_context):
    resumen = (continuity_context or {}).get('resumen', {})
    if not resumen:
        return

    sheet = workbook.create_sheet('Resumen')
    sheet.append(['Metrica', 'Valor'])
    header_fill = PatternFill(start_color='27348B', end_color='27348B', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    rows = [
        ('total_principal', resumen.get('total_documentos_principal', 0)),
        ('total_secundaria', resumen.get('total_documentos_secundaria', 0)),
        ('total_consolidado', resumen.get('total_documentos_consolidado', 0)),
        ('total_recuperados_secundaria', resumen.get('total_documentos_recuperados_secundaria', 0)),
        ('total_duplicados_entre_carpetas', resumen.get('total_duplicados_entre_carpetas', 0)),
        ('total_faltantes_pel_base', resumen.get('total_faltantes_detectados', 0)),
        ('total_faltantes_subparte', resumen.get('total_faltantes_subparte', 0)),
        ('total_faltantes_tomo', resumen.get('total_faltantes_tomo', 0)),
        ('total_faltantes_no_encontrados', resumen.get('total_faltantes_no_encontrados', 0)),
        ('total_cambios_de_rango', resumen.get('total_cambios_de_rango', 0)),
        ('total_omitidos_sistema', resumen.get('total_omitidos_sistema', 0)),
        ('total_omitidos_temporales', resumen.get('total_omitidos_temporales', 0)),
        ('total_omitidos_ocultos', resumen.get('total_omitidos_ocultos', 0)),
        ('total_carpetas_omitidas', resumen.get('total_carpetas_omitidas', 0)),
        ('total_no_soportados', resumen.get('total_no_soportados', 0)),
        ('carpeta_secundaria_usada', resumen.get('carpeta_secundaria_usada', False)),
        ('inventario_secundaria_completo', resumen.get('inventario_secundaria_completo', False)),
        ('carpeta_secundaria', resumen.get('carpeta_secundaria', '')),
    ]
    for row in rows:
        sheet.append(row)

    sheet.column_dimensions['A'].width = 38
    sheet.column_dimensions['B'].width = 80


def _normalize_path(path):
    return str(path or '').replace('\\', '/').lower()


def _continuity_blocked_bases(context):
    blocked = {}
    if not context:
        return blocked
    for item in context.get('subpartes', {}).get('faltantes', []):
        base = _safe_int(item.get('pel_base'))
        if base:
            blocked[base] = f"Falta subparte documental: {item.get('consecutivo_faltante')}"
    for item in context.get('tomos', {}).get('faltantes', []):
        base = _safe_int(item.get('pel_base'))
        if base:
            blocked[base] = f"Falta tomo o parte documental: {item.get('consecutivo_faltante')}"
    return blocked


def _safe_int(value):
    if value in (None, ''):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _append_observation(current, extra):
    if not current:
        return extra
    return f'{current}; {extra}'


def _format_missing_pel_range(start, end):
    if start == end:
        return f'PEL {start}'
    if end - start <= 10:
        return ', '.join(f'PEL {number}' for number in range(start, end + 1))
    return f'PEL {start}..PEL {end}'


