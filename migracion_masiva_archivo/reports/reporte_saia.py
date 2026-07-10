"""
Reporte Excel de PEL cargados a SAIA.

Formato:
  - Fila 1: "PEL 2019 CARGADOS A SAIA" (merge A1:J1, fondo azul, blanco negrilla)
  - Fila 2: "REPORTE DEL DD DE MES YYYY" (merge A2:J2, mismo estilo)
  - Filas 3+: consecutivos SAIA ordenados, distribuidos en 10 columnas de izquierda a derecha
  - Hoja DETALLE opcional con todos los campos tecnicos
"""

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from migracion_masiva_archivo.models import IntentoCargaSAIA, LogProcesoDocumental


MESES_ES = {
    1: 'ENERO', 2: 'FEBRERO', 3: 'MARZO', 4: 'ABRIL',
    5: 'MAYO', 6: 'JUNIO', 7: 'JULIO', 8: 'AGOSTO',
    9: 'SEPTIEMBRE', 10: 'OCTUBRE', 11: 'NOVIEMBRE', 12: 'DICIEMBRE',
}

_HEADER_FILL = PatternFill(start_color='27348B', end_color='27348B', fill_type='solid')
_HEADER_FONT = Font(color='FFFFFF', bold=True, size=11)
_HEADER_ALIGN = Alignment(horizontal='center', vertical='center')
_BODY_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
_THIN = Side(style='thin')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_DETALLE_HEADERS = [
    'lote', 'documento_id', 'nombre_archivo', 'ruta_archivo',
    'consecutivo', 'asunto_saia', 'estado_proceso', 'intento_id',
    'fecha_hora', 'mensaje_saia', 'id_documento_saia', 'screenshot', 'error',
]


# ---------------------------------------------------------------------------
# Consultas
# ---------------------------------------------------------------------------

def get_asuntos_exitosos_por_fecha(fecha):
    """Retorna lista de {asunto_saia, intento} unicos cargados en la fecha dada."""
    intentos = (
        IntentoCargaSAIA.objects
        .filter(exitoso=True, creado__date=fecha)
        .select_related('documento', 'documento__metadata', 'documento__lote')
        .order_by('creado')
    )
    return _dedup_asuntos(intentos)


def get_asuntos_exitosos_por_lote(lote_id):
    """Retorna lista de {asunto_saia, intento} unicos cargados en el lote dado."""
    intentos = (
        IntentoCargaSAIA.objects
        .filter(exitoso=True, documento__lote_id=lote_id)
        .select_related('documento', 'documento__metadata', 'documento__lote')
        .order_by('creado')
    )
    return _dedup_asuntos(intentos)


def _dedup_asuntos(intentos):
    seen = set()
    result = []
    for intento in intentos:
        metadata = intento.request_metadata or {}
        asunto = metadata.get('asunto_saia', '')
        if asunto and asunto not in seen:
            seen.add(asunto)
            result.append({'asunto_saia': asunto, 'intento': intento})
    return result


# ---------------------------------------------------------------------------
# Construccion del Excel
# ---------------------------------------------------------------------------

def build_reporte_saia_xlsx(items, fecha_reporte, con_detalle=False):
    """
    Construye el Excel de PEL cargados a SAIA.

    items        : lista de dicts {'asunto_saia': str, 'intento': IntentoCargaSAIA}
    fecha_reporte: datetime.date
    con_detalle  : agrega hoja DETALLE con campos tecnicos
    """
    wb = Workbook()
    ws = wb.active
    ws.title = 'PEL CARGADOS'
    _build_main_sheet(ws, items, fecha_reporte)

    if con_detalle and items:
        wd = wb.create_sheet('DETALLE')
        _build_detalle_sheet(wd, items)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def _build_main_sheet(ws, items, fecha_reporte):
    mes = MESES_ES[fecha_reporte.month]
    titulo = 'PEL 2019 CARGADOS A SAIA'
    subtitulo = f'REPORTE DEL {fecha_reporte.day} DE {mes} {fecha_reporte.year}'

    ws.merge_cells('A1:J1')
    _apply_header(ws['A1'], titulo)
    ws.row_dimensions[1].height = 24

    ws.merge_cells('A2:J2')
    _apply_header(ws['A2'], subtitulo)
    ws.row_dimensions[2].height = 24

    for col in range(1, 11):
        ws.column_dimensions[get_column_letter(col)].width = 22

    asuntos = sorted(
        [item['asunto_saia'] for item in items],
        key=_sort_key_pel,
    )

    row = 3
    col = 1
    for asunto in asuntos:
        cell = ws.cell(row=row, column=col, value=asunto)
        cell.border = _BORDER
        cell.alignment = _BODY_ALIGN
        col += 1
        if col > 10:
            col = 1
            row += 1

    # Rellena celdas vacias al final de la ultima fila con borde
    if col > 1:
        while col <= 10:
            cell = ws.cell(row=row, column=col, value='')
            cell.border = _BORDER
            col += 1


def _build_detalle_sheet(wd, items):
    wd.append(_DETALLE_HEADERS)
    for cell in wd[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGN

    for item in items:
        intento = item['intento']
        doc = intento.documento
        meta = getattr(doc, 'metadata', None)
        respuesta = intento.respuesta or {}
        wd.append([
            doc.lote_id,
            doc.id,
            doc.nombre_archivo,
            doc.ruta_archivo,
            meta.consecutivo if meta else '',
            item['asunto_saia'],
            doc.estado_proceso,
            intento.id,
            str(intento.creado),
            intento.mensaje_error or '',
            intento.id_documento_saia or '',
            respuesta.get('screenshot', ''),
            intento.mensaje_error or '',
        ])

    for column_cells in wd.columns:
        max_len = max(len(str(c.value or '')) for c in column_cells)
        wd.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 12), 60)


def _apply_header(cell, value):
    cell.value = value
    cell.fill = _HEADER_FILL
    cell.font = _HEADER_FONT
    cell.alignment = _HEADER_ALIGN


# ---------------------------------------------------------------------------
# Ordenamiento de asuntos PEL
# ---------------------------------------------------------------------------

def _sort_key_pel(asunto):
    """
    Orden ascendente para consecutivos PEL:
      PEL N          -> (N, 0, 0, 0, 0)
      PEL N -S       -> (N, 0, 0, 1, S)
      PEL N TOMO T-S -> (N, 1, T, 1, S)
    """
    tomo_m = re.match(r'PEL (\d+) TOMO (\d+)-(\d+)', asunto)
    if tomo_m:
        return (int(tomo_m.group(1)), 1, int(tomo_m.group(2)), 1, int(tomo_m.group(3)))

    suffix_m = re.match(r'PEL (\d+) -(\d+)', asunto)
    if suffix_m:
        return (int(suffix_m.group(1)), 0, 0, 1, int(suffix_m.group(2)))

    simple_m = re.match(r'PEL (\d+)', asunto)
    if simple_m:
        return (int(simple_m.group(1)), 0, 0, 0, 0)

    return (999999999, 0, 0, 0, 0)


# ---------------------------------------------------------------------------
# Envio por correo
# ---------------------------------------------------------------------------

def send_reporte_saia_email(
    xlsx_bytes, fecha_reporte, destinatarios, lote=None,
    xlsx_fase5_bytes=None, pendientes=0,
):
    """
    Envia el reporte Excel como adjunto.
    Si xlsx_fase5_bytes esta presente, adjunta tambien el reporte detallado Fase 5.
    Si pendientes > 0, el asunto y cuerpo indican que el reporte es parcial.
    Registra el resultado en LogProcesoDocumental (si se pasa lote).
    No guarda credenciales: usa la configuracion EMAIL_* de settings.py.
    """
    import re as _re
    from django.conf import settings
    from django.core.mail import EmailMessage

    if not xlsx_bytes:
        raise ValueError('El reporte Excel esta vacio, no se puede enviar.')

    mes = MESES_ES[fecha_reporte.month]
    fecha_str = f'{fecha_reporte.day} DE {mes} {fecha_reporte.year}'
    fecha_safe = fecha_reporte.strftime('%Y%m%d')
    mime_xlsx = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    nombre_lote = getattr(lote, 'nombre', '') or 'LOTE'
    nombre_lote_safe = _re.sub(r'[^\w\s-]', '', nombre_lote).strip().replace(' ', '_')
    nombre_xlsx = f'Reporte_SAIA_{nombre_lote_safe}_{fecha_safe}.xlsx'

    if pendientes:
        subject = f'{nombre_lote} CARGADOS A SAIA - REPORTE PARCIAL DEL {fecha_str} ({pendientes} doc(s) pendientes)'
        aviso_pendientes = (
            f'\nATENCION: Este es un reporte PARCIAL. Quedaron {pendientes} documento(s) '
            f'listos sin cargar por el limite de documentos por ciclo. '
            f'Vuelva a ejecutar el proceso para cargar el siguiente lote.\n'
        )
    else:
        subject = f'{nombre_lote} CARGADOS A SAIA - REPORTE DEL {fecha_str}'
        aviso_pendientes = ''

    email = EmailMessage(
        subject=subject,
        body=(
            f'Estimados,\n\n'
            f'Se adjunta el reporte de {nombre_lote} cargados exitosamente a SAIA '
            f'del {fecha_str}.\n'
            f'{aviso_pendientes}\n'
            f'Este mensaje fue generado automaticamente por el sistema de gestion documental.'
        ),
        from_email=settings.DEFAULT_FROM_EMAIL,
        to=destinatarios,
    )
    email.attach(nombre_xlsx, xlsx_bytes, mime_xlsx)

    if xlsx_fase5_bytes:
        nombre_fase5 = f'Detalle_SAIA_{nombre_lote_safe}_{fecha_safe}.xlsx'
        email.attach(nombre_fase5, xlsx_fase5_bytes, mime_xlsx)

    try:
        email.send(fail_silently=False)
        _registrar_correo(lote, 'INFO', 'reporte_saia_correo_enviado',
                          f'Reporte enviado a: {", ".join(destinatarios)}',
                          {'destinatarios': destinatarios, 'archivo': nombre_xlsx})
        return {'enviado': True, 'destinatarios': destinatarios, 'archivo': nombre_xlsx}
    except Exception as exc:
        _registrar_correo(lote, 'ERROR', 'reporte_saia_correo_error',
                          f'Error al enviar reporte: {exc}',
                          {'error': str(exc), 'destinatarios': destinatarios})
        raise


def _registrar_correo(lote, nivel, evento, mensaje, detalle):
    if lote is None:
        return
    LogProcesoDocumental.objects.create(
        lote=lote,
        nivel=nivel,
        evento=evento,
        mensaje=mensaje,
        detalle=detalle or {},
    )


